"""Machines (main machines / production lines) and their components.

Hierarchy: machine → components → faults. A fault always names its machine (`Kvar.oprema_id`)
and, when the affected component is known, that component (`Kvar.komponenta_id`, which must
belong to the same machine). Machine-level counts therefore include component faults exactly once.
"""
import logging
from io import BytesIO
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import case, func, or_
from sqlalchemy.orm import Session, joinedload

from ..auth import get_current_user, require_admin, require_voditelj_or_admin, role_value
from ..config import settings
from ..database import get_db
from ..labels import Label, asset_link, labels_pdf
from ..models import (
    KomponentaOpreme,
    Korisnik,
    Kvar,
    KvarRezervniDio,
    Oprema,
    PrioritetEnum,
    StatusKvaraEnum,
    StatusOpremaEnum,
    TvrtkaEnum,
    nova_qr_oznaka,
)
from ..schemas import (
    KomponentaCreate,
    KomponentaDetalj,
    KomponentaOut,
    KomponentaUpdate,
    OpremaCreate,
    OpremaOut,
    OpremaUpdate,
    OpremaWithStats,
)

router = APIRouter(prefix="/api/oprema", tags=["oprema"])
logger = logging.getLogger("adixpaper")

OPEN_STATUSES = [StatusKvaraEnum.novo, StatusKvaraEnum.u_rjesavanju, StatusKvaraEnum.na_cekanju]
MACHINE_NOT_FOUND = "Oprema nije pronađena"
COMPONENT_NOT_FOUND = "Komponenta nije pronađena"


def _norm(value: Optional[str]) -> str:
    """Case- and spacing-insensitive form for duplicate checks ("Odmotač  1" = "ODMOTAČ 1")."""
    return " ".join((value or "").split()).casefold()


def _check_identifiers(
    db: Session,
    *,
    inventurni_broj: Optional[str] = None,
    serijski_broj: Optional[str] = None,
    proizvodac: Optional[str] = None,
    machine_id: Optional[int] = None,
    component_id: Optional[int] = None,
    is_component: bool,
) -> None:
    """Business rules for identifiers, checked only when a value is entered (both stay optional):

    - an inventory number identifies one asset in the plant: unique across machines AND components;
    - a machine's serial number is unique among machines (existing database constraint);
    - a component's serial number is unique per manufacturer (different makers may reuse numbers).
    Comparisons ignore case and surrounding spaces.
    """
    if inventurni_broj:
        wanted = inventurni_broj.strip().lower()
        machine_clash = db.query(Oprema.id).filter(func.lower(Oprema.inventurni_broj) == wanted)
        if machine_id is not None:
            machine_clash = machine_clash.filter(Oprema.id != machine_id)
        component_clash = db.query(KomponentaOpreme.id).filter(func.lower(KomponentaOpreme.inventurni_broj) == wanted)
        if component_id is not None:
            component_clash = component_clash.filter(KomponentaOpreme.id != component_id)
        if machine_clash.first() or component_clash.first():
            raise HTTPException(status_code=400, detail="Inventurni broj već postoji")
    if serijski_broj:
        wanted = serijski_broj.strip().lower()
        if not is_component:
            clash = db.query(Oprema.id).filter(func.lower(Oprema.serijski_broj) == wanted)
            if machine_id is not None:
                clash = clash.filter(Oprema.id != machine_id)
            if clash.first():
                raise HTTPException(status_code=400, detail="Serijski broj već postoji")
        else:
            clash = db.query(KomponentaOpreme).filter(func.lower(KomponentaOpreme.serijski_broj) == wanted)
            if component_id is not None:
                clash = clash.filter(KomponentaOpreme.id != component_id)
            if any(_norm(other.proizvodac) == _norm(proizvodac) for other in clash.all()):
                raise HTTPException(status_code=400, detail="Serijski broj već postoji za tog proizvođača")


def _check_component_name(db: Session, oprema_id: int, naziv: str, component_id: Optional[int] = None) -> None:
    """One component of a given name per machine, so workers and labels can tell them apart."""
    siblings = db.query(KomponentaOpreme.id, KomponentaOpreme.naziv).filter(KomponentaOpreme.oprema_id == oprema_id)
    if any(_norm(name) == _norm(naziv) and row_id != component_id for row_id, name in siblings.all()):
        raise HTTPException(status_code=400, detail="Komponenta s tim nazivom već postoji na ovoj mašini")


def _parts_cost(db: Session, oprema_id: Optional[int] = None) -> dict:
    """Spare-part cost per machine at the prices recorded when the parts were used."""
    query = db.query(
        Kvar.oprema_id,
        func.coalesce(func.sum(KvarRezervniDio.kolicina * func.coalesce(KvarRezervniDio.jedinicna_cijena, 0)), 0),
    ).join(Kvar, KvarRezervniDio.kvar_id == Kvar.id)
    if oprema_id is not None:
        query = query.filter(Kvar.oprema_id == oprema_id)
    return {machine: float(total or 0) for machine, total in query.group_by(Kvar.oprema_id).all()}


def _machine_stats(db: Session, oprema_id: Optional[int] = None) -> dict:
    """Per machine: all faults, open faults, open urgent faults (component faults included once),
    labour cost, last report, active components."""
    open_case = func.sum(case((Kvar.status.in_(OPEN_STATUSES), 1), else_=0))
    urgent_case = func.sum(case((Kvar.status.in_(OPEN_STATUSES) & (Kvar.prioritet == PrioritetEnum.hitno), 1), else_=0))
    faults = db.query(
        Kvar.oprema_id,
        func.count(Kvar.id).label("broj"),
        open_case.label("otvoreni"),
        urgent_case.label("hitni"),
        func.coalesce(func.sum(Kvar.troskovi), 0).label("troskovi"),
        func.max(Kvar.datum_prijave).label("zadnji"),
    )
    components = db.query(KomponentaOpreme.oprema_id, func.count(KomponentaOpreme.id)).filter(KomponentaOpreme.aktivna == True)  # noqa: E712
    if oprema_id is not None:
        faults = faults.filter(Kvar.oprema_id == oprema_id)
        components = components.filter(KomponentaOpreme.oprema_id == oprema_id)
    return {
        "faults": {row.oprema_id: row for row in faults.group_by(Kvar.oprema_id).all()},
        "components": dict(components.group_by(KomponentaOpreme.oprema_id).all()),
        "parts": _parts_cost(db, oprema_id),
    }


def _with_stats(o: Oprema, stats: dict, user: Korisnik) -> OpremaWithStats:
    item = OpremaWithStats.model_validate(o)
    row = stats["faults"].get(o.id)
    if row:
        item.broj_kvarova = row.broj or 0
        item.aktivni_kvarovi = int(row.otvoreni or 0)
        item.hitni_kvarovi = int(row.hitni or 0)
        item.ukupni_troskovi = float(row.troskovi or 0) + stats["parts"].get(o.id, 0.0)
        item.zadnji_kvar = row.zadnji
    item.broj_komponenti = stats["components"].get(o.id, 0)
    return _for_user(item, user)


def _for_user(item: OpremaWithStats, user: Korisnik) -> OpremaWithStats:
    """Workers pick machines to report on; repair costs are not theirs to see."""
    if role_value(user.rola) == "radnik":
        item.ukupni_troskovi = None
    return item


def _component_counts(db: Session, component_ids: list) -> dict:
    if not component_ids:
        return {}
    rows = db.query(
        Kvar.komponenta_id,
        func.count(Kvar.id),
        func.sum(case((Kvar.status.in_(OPEN_STATUSES), 1), else_=0)),
    ).filter(Kvar.komponenta_id.in_(component_ids)).group_by(Kvar.komponenta_id).all()
    return {component_id: (total or 0, int(open_ or 0)) for component_id, total, open_ in rows}


def _component_out(component: KomponentaOpreme, counts: dict, model=KomponentaOut):
    item = model.model_validate(component)
    item.ukupno_kvarova, item.aktivni_kvarovi = counts.get(component.id, (0, 0))
    return item


@router.get("", response_model=List[OpremaWithStats])
def get_oprema(
    search: Optional[str] = Query(None, max_length=200),
    status: Optional[StatusOpremaEnum] = Query(None),
    tvrtka: Optional[TvrtkaEnum] = Query(None),
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(get_current_user),
):
    """Machines with counts. `search` matches the machine's own fields and its components'
    (name, model, manufacturer, inventory and serial numbers)."""
    query = db.query(Oprema)
    if search and search.strip():
        pattern = f"%{search.strip()}%"
        component_match = db.query(KomponentaOpreme.id).filter(
            KomponentaOpreme.oprema_id == Oprema.id,
            or_(
                KomponentaOpreme.naziv.ilike(pattern),
                KomponentaOpreme.model.ilike(pattern),
                KomponentaOpreme.proizvodac.ilike(pattern),
                KomponentaOpreme.inventurni_broj.ilike(pattern),
                KomponentaOpreme.serijski_broj.ilike(pattern),
            ),
        ).exists()
        query = query.filter(or_(
            Oprema.naziv.ilike(pattern),
            Oprema.tip.ilike(pattern),
            Oprema.serijski_broj.ilike(pattern),
            Oprema.inventurni_broj.ilike(pattern),
            Oprema.lokacija.ilike(pattern),
            Oprema.proizvodac.ilike(pattern),
            Oprema.model.ilike(pattern),
            component_match,
        ))
    if status:
        query = query.filter(Oprema.status == status)
    if tvrtka:
        query = query.filter(Oprema.tvrtka == tvrtka)
    stats = _machine_stats(db)
    return [_with_stats(o, stats, current_user) for o in query.order_by(Oprema.naziv).all()]


def machine_link(base_url: str, oprema_id: int) -> str:
    """The address on labels printed before 2026-10-30 (still served by the app's /m/<id> route)."""
    return f"{base_url.rstrip('/')}/m/{oprema_id}"


def _base_url(request: Request) -> str:
    return settings.PUBLIC_URL or str(request.base_url)


def _machine_label(machine: Oprema, base_url: str) -> Label:
    return Label(
        url=asset_link(base_url, machine.qr_oznaka),
        title=machine.naziv,
        meta=" · ".join(part for part in (machine.model, machine.proizvodac, machine.lokacija) if part),
        kind="Mašina",
    )


def _component_label(component: KomponentaOpreme, machine: Oprema, base_url: str) -> Label:
    return Label(
        url=asset_link(base_url, component.qr_oznaka),
        title=component.naziv,
        parent=machine.naziv,
        meta=" · ".join(part for part in (component.model, component.proizvodac) if part),
        kind="Komponenta",
    )


def _pdf_response(output: BytesIO, filename: str) -> StreamingResponse:
    return StreamingResponse(output, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _ensure_token(db: Session, asset) -> None:
    if not asset.qr_oznaka:
        asset.qr_oznaka = nova_qr_oznaka()
        db.commit()


@router.get("/qr-oznake.pdf")
def machine_qr_labels(
    request: Request,
    ids: Optional[str] = Query(None, description="Comma-separated machine ids; default: every machine not decommissioned"),
    db: Session = Depends(get_db),
    _: Korisnik = Depends(require_voditelj_or_admin),
):
    """A4 sheet of cut-out labels, one per machine (the machines' own labels, not their components')."""
    query = db.query(Oprema)
    if ids:
        try:
            wanted = [int(value) for value in ids.split(",") if value.strip()]
        except ValueError:
            raise HTTPException(status_code=422, detail="Neispravna vrijednost filtra") from None
        query = query.filter(Oprema.id.in_(wanted))
    else:
        query = query.filter(Oprema.status != StatusOpremaEnum.raspisano)
    machines = query.order_by(Oprema.naziv).all()
    if not machines:
        raise HTTPException(status_code=404, detail=MACHINE_NOT_FOUND)
    for machine in machines:
        _ensure_token(db, machine)
    base_url = _base_url(request)
    return _pdf_response(labels_pdf([_machine_label(m, base_url) for m in machines], "Adix Paper – QR oznake mašina"), "adix-paper-qr-oznake.pdf")


@router.get("/export.xlsx")
def export_assets(db: Session = Depends(get_db), _: Korisnik = Depends(require_voditelj_or_admin)):
    """Register of machines and their components, with identifiers and open faults."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    stats = _machine_stats(db)
    machines = db.query(Oprema).options(joinedload(Oprema.komponente)).order_by(Oprema.naziv).all()
    counts = _component_counts(db, [c.id for m in machines for c in m.komponente])
    wb = Workbook()
    ws = wb.active
    ws.title = "Mašine i komponente"
    ws.append(["Vrsta", "Mašina", "Komponenta", "Proizvođač", "Model", "Inventurni broj", "Serijski broj", "Aktivna", "QR oznaka", "Otvoreni kvarovi"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    for machine in machines:
        row = stats["faults"].get(machine.id)
        ws.append([
            "Mašina", machine.naziv, "", machine.proizvodac or "", machine.model or "", machine.inventurni_broj or "",
            machine.serijski_broj or "", "da" if machine.aktivna else "ne", machine.qr_oznaka or "", int(row.otvoreni or 0) if row else 0,
        ])
        for component in sorted(machine.komponente, key=lambda c: _norm(c.naziv)):
            ws.append([
                "Komponenta", machine.naziv, component.naziv, component.proizvodac or "", component.model or "",
                component.inventurni_broj or "", component.serijski_broj or "", "da" if component.aktivna else "ne",
                component.qr_oznaka or "", counts.get(component.id, (0, 0))[1],
            ])
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 10), 42)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="adix-paper-masine-komponente.xlsx"'},
    )


@router.get("/komponente/trazi", response_model=List[KomponentaDetalj])
def search_components(
    q: str = Query(..., min_length=1, max_length=100),
    db: Session = Depends(get_db),
    _: Korisnik = Depends(require_voditelj_or_admin),
):
    """Components of any machine by name, model, manufacturer, inventory or serial number."""
    pattern = f"%{q.strip()}%"
    components = db.query(KomponentaOpreme).options(joinedload(KomponentaOpreme.oprema)).filter(or_(
        KomponentaOpreme.naziv.ilike(pattern),
        KomponentaOpreme.model.ilike(pattern),
        KomponentaOpreme.proizvodac.ilike(pattern),
        KomponentaOpreme.inventurni_broj.ilike(pattern),
        KomponentaOpreme.serijski_broj.ilike(pattern),
    )).order_by(KomponentaOpreme.naziv).limit(30).all()
    counts = _component_counts(db, [c.id for c in components])
    return [_component_out(c, counts, KomponentaDetalj) for c in components]


@router.get("/komponente/{komponenta_id}", response_model=KomponentaDetalj)
def get_komponenta(komponenta_id: int, db: Session = Depends(get_db), _: Korisnik = Depends(require_voditelj_or_admin)):
    component = db.query(KomponentaOpreme).options(joinedload(KomponentaOpreme.oprema)).filter(KomponentaOpreme.id == komponenta_id).first()
    if not component:
        raise HTTPException(status_code=404, detail=COMPONENT_NOT_FOUND)
    return _component_out(component, _component_counts(db, [component.id]), KomponentaDetalj)


@router.get("/komponente/{komponenta_id}/qr-oznaka.pdf")
def component_qr_label(request: Request, komponenta_id: int, db: Session = Depends(get_db), _: Korisnik = Depends(require_voditelj_or_admin)):
    component = db.query(KomponentaOpreme).options(joinedload(KomponentaOpreme.oprema)).filter(KomponentaOpreme.id == komponenta_id).first()
    if not component:
        raise HTTPException(status_code=404, detail=COMPONENT_NOT_FOUND)
    _ensure_token(db, component)
    output = labels_pdf([_component_label(component, component.oprema, _base_url(request))], f"Adix Paper – {component.naziv}")
    return _pdf_response(output, "adix-paper-qr-komponenta.pdf")


@router.post("/komponente/{komponenta_id}/nova-qr-oznaka", response_model=KomponentaOut)
def regenerate_component_token(komponenta_id: int, db: Session = Depends(get_db), user: Korisnik = Depends(require_voditelj_or_admin)):
    """A new QR identifier: labels printed with the old one stop working. Only on explicit request
    (e.g. a label was copied where it should not be)."""
    component = db.query(KomponentaOpreme).filter(KomponentaOpreme.id == komponenta_id).first()
    if not component:
        raise HTTPException(status_code=404, detail=COMPONENT_NOT_FOUND)
    component.qr_oznaka = nova_qr_oznaka()
    db.commit()
    db.refresh(component)
    logger.warning("New QR identifier for component %s by %s; old labels no longer work", component.id, user.username)
    return _component_out(component, _component_counts(db, [component.id]))


@router.get("/{oprema_id}", response_model=OpremaWithStats)
def get_oprema_by_id(oprema_id: int, db: Session = Depends(get_db), current_user: Korisnik = Depends(get_current_user)):
    o = db.query(Oprema).filter(Oprema.id == oprema_id).first()
    if not o:
        raise HTTPException(status_code=404, detail=MACHINE_NOT_FOUND)
    return _with_stats(o, _machine_stats(db, o.id), current_user)


@router.get("/{oprema_id}/qr-oznake.pdf")
def machine_and_component_labels(
    request: Request,
    oprema_id: int,
    komponente: Optional[str] = Query(None, description="Comma-separated component ids; default: every active component"),
    masina: bool = Query(True, description="Include the machine's own label"),
    db: Session = Depends(get_db),
    _: Korisnik = Depends(require_voditelj_or_admin),
):
    """One sheet for a production line: its own label, then one per component."""
    machine = db.query(Oprema).options(joinedload(Oprema.komponente)).filter(Oprema.id == oprema_id).first()
    if not machine:
        raise HTTPException(status_code=404, detail=MACHINE_NOT_FOUND)
    if komponente:
        try:
            wanted = {int(value) for value in komponente.split(",") if value.strip()}
        except ValueError:
            raise HTTPException(status_code=422, detail="Neispravna vrijednost filtra") from None
        chosen = [c for c in machine.komponente if c.id in wanted]
        if len(chosen) != len(wanted):
            raise HTTPException(status_code=404, detail="Komponenta nije pronađena za odabranu opremu")
    else:
        chosen = [c for c in machine.komponente if c.aktivna]
    chosen.sort(key=lambda c: _norm(c.naziv))
    for asset in [machine, *chosen]:
        _ensure_token(db, asset)
    base_url = _base_url(request)
    labels = ([_machine_label(machine, base_url)] if masina else []) + [_component_label(c, machine, base_url) for c in chosen]
    if not labels:
        raise HTTPException(status_code=404, detail=COMPONENT_NOT_FOUND)
    return _pdf_response(labels_pdf(labels, f"Adix Paper – {machine.naziv}"), "adix-paper-qr-oznake-masine.pdf")


@router.post("/{oprema_id}/nova-qr-oznaka", response_model=OpremaOut)
def regenerate_machine_token(oprema_id: int, db: Session = Depends(get_db), user: Korisnik = Depends(require_voditelj_or_admin)):
    machine = db.query(Oprema).filter(Oprema.id == oprema_id).first()
    if not machine:
        raise HTTPException(status_code=404, detail=MACHINE_NOT_FOUND)
    machine.qr_oznaka = nova_qr_oznaka()
    db.commit()
    db.refresh(machine)
    logger.warning("New QR identifier for machine %s by %s; old /q/ labels no longer work", machine.id, user.username)
    return machine


@router.post("", response_model=OpremaOut)
def create_oprema(data: OpremaCreate, db: Session = Depends(get_db), _: Korisnik = Depends(require_voditelj_or_admin)):
    _check_identifiers(db, inventurni_broj=data.inventurni_broj, serijski_broj=data.serijski_broj, is_component=False)
    oprema = Oprema(**data.model_dump())
    db.add(oprema)
    db.commit()
    db.refresh(oprema)
    return oprema


@router.put("/{oprema_id}", response_model=OpremaOut)
def update_oprema(oprema_id: int, data: OpremaUpdate, db: Session = Depends(get_db), _: Korisnik = Depends(require_voditelj_or_admin)):
    oprema = db.query(Oprema).filter(Oprema.id == oprema_id).first()
    if not oprema:
        raise HTTPException(status_code=404, detail=MACHINE_NOT_FOUND)
    update_data = data.model_dump(exclude_unset=True)
    _check_identifiers(
        db, inventurni_broj=update_data.get("inventurni_broj"), serijski_broj=update_data.get("serijski_broj"),
        machine_id=oprema_id, is_component=False,
    )
    for field, value in update_data.items():
        setattr(oprema, field, value)
    db.commit()
    db.refresh(oprema)
    return oprema


@router.get("/{oprema_id}/komponente", response_model=List[KomponentaOut])
def get_komponente(oprema_id: int, db: Session = Depends(get_db), current_user: Korisnik = Depends(get_current_user)):
    """A machine's components (everyone: workers choose one when reporting). Fault counts for leads."""
    if not db.query(Oprema.id).filter(Oprema.id == oprema_id).first():
        raise HTTPException(status_code=404, detail=MACHINE_NOT_FOUND)
    components = db.query(KomponentaOpreme).filter(KomponentaOpreme.oprema_id == oprema_id).order_by(KomponentaOpreme.naziv).all()
    counts = _component_counts(db, [c.id for c in components]) if role_value(current_user.rola) != "radnik" else {}
    return [_component_out(c, counts) for c in components]


@router.post("/{oprema_id}/komponente", response_model=KomponentaOut, status_code=201)
def create_komponenta(
    oprema_id: int,
    data: KomponentaCreate,
    db: Session = Depends(get_db),
    _: Korisnik = Depends(require_voditelj_or_admin),
):
    if not db.query(Oprema.id).filter(Oprema.id == oprema_id).first():
        raise HTTPException(status_code=404, detail=MACHINE_NOT_FOUND)
    naziv = " ".join(data.naziv.split())
    if not naziv:
        raise HTTPException(status_code=400, detail="Naziv komponente je obavezan")
    _check_component_name(db, oprema_id, naziv)
    _check_identifiers(db, inventurni_broj=data.inventurni_broj, serijski_broj=data.serijski_broj, proizvodac=data.proizvodac, is_component=True)
    komponenta = KomponentaOpreme(
        oprema_id=oprema_id,
        naziv=naziv,
        opis=data.opis,
        proizvodac=data.proizvodac,
        model=data.model,
        inventurni_broj=data.inventurni_broj,
        serijski_broj=data.serijski_broj,
        aktivna=data.aktivna,
    )
    db.add(komponenta)
    db.commit()
    db.refresh(komponenta)
    return komponenta


@router.put("/komponente/{komponenta_id}", response_model=KomponentaOut)
def update_komponenta(
    komponenta_id: int,
    data: KomponentaUpdate,
    db: Session = Depends(get_db),
    _: Korisnik = Depends(require_voditelj_or_admin),
):
    komponenta = db.query(KomponentaOpreme).filter(KomponentaOpreme.id == komponenta_id).first()
    if not komponenta:
        raise HTTPException(status_code=404, detail=COMPONENT_NOT_FOUND)
    update_data = data.model_dump(exclude_unset=True)
    if update_data.get("naziv") is not None:
        update_data["naziv"] = " ".join(update_data["naziv"].split())
        if not update_data["naziv"]:
            raise HTTPException(status_code=400, detail="Naziv komponente je obavezan")
        _check_component_name(db, komponenta.oprema_id, update_data["naziv"], komponenta.id)
    _check_identifiers(
        db,
        inventurni_broj=update_data.get("inventurni_broj"),
        serijski_broj=update_data.get("serijski_broj", komponenta.serijski_broj if "proizvodac" in update_data else None),
        proizvodac=update_data.get("proizvodac", komponenta.proizvodac),
        component_id=komponenta.id,
        is_component=True,
    )
    for field, value in update_data.items():
        setattr(komponenta, field, value)
    db.commit()
    db.refresh(komponenta)
    return _component_out(komponenta, _component_counts(db, [komponenta.id]))


@router.delete("/komponente/{komponenta_id}", status_code=204)
def delete_komponenta(
    komponenta_id: int,
    db: Session = Depends(get_db),
    _: Korisnik = Depends(require_voditelj_or_admin),
):
    """Deletes a component that never had a fault. One with fault history is kept for the record:
    the request is refused and the component can be deactivated instead."""
    komponenta = db.query(KomponentaOpreme).filter(KomponentaOpreme.id == komponenta_id).first()
    if not komponenta:
        raise HTTPException(status_code=404, detail=COMPONENT_NOT_FOUND)
    if db.query(Kvar.id).filter(Kvar.komponenta_id == komponenta_id).first():
        raise HTTPException(status_code=409, detail="Komponenta ima evidentirane kvarove i ne može se obrisati. Deaktivirajte je.")
    db.delete(komponenta)
    db.commit()


@router.delete("/{oprema_id}")
def delete_oprema(oprema_id: int, db: Session = Depends(get_db), _: Korisnik = Depends(require_admin)):
    oprema = db.query(Oprema).filter(Oprema.id == oprema_id).first()
    if not oprema:
        raise HTTPException(status_code=404, detail=MACHINE_NOT_FOUND)
    if db.query(Kvar.id).filter(Kvar.oprema_id == oprema_id).first():
        raise HTTPException(
            status_code=409,
            detail="Mašina ima evidentirane kvarove i ne može se obrisati. Postavite status na Raspisano.",
        )
    db.delete(oprema)
    db.commit()
    return {"detail": "Oprema obrisana"}
