from fastapi import APIRouter, Depends, File, HTTPException, BackgroundTasks, Query, Response, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import desc, func, or_
from sqlalchemy.exc import IntegrityError
from typing import List, Optional
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
import json
from urllib.parse import urlsplit
from types import SimpleNamespace
import logging

from ..database import get_db, SessionLocal
from ..models import (
    Kvar,
    Oprema,
    Korisnik,
    PushPretplata,
    KvarDogadjaj,
    KvarPrilog,
    KvarRezervniDio,
    KomponentaOpreme,
    RezervniDio,
    RolaEnum,
    PrioritetEnum,
    StatusKvaraEnum,
    SLA_SATI,
    RAZLOZI_CEKANJA,
)
from ..schemas import (
    KVAR_INTERNAL_FIELDS,
    KvarCreate,
    KvarDogadjajOut,
    KvarKomentarCreate,
    KvarListOut,
    KvarNaMasiniOut,
    KvarPrilogOut,
    KvarNovoOut,
    KvarOut,
    KvarVoditeljUpdate,
)
from ..auth import get_current_user, require_admin, require_voditelj_or_admin, role_value
from ..timeutil import utcnow, as_utc, local_naive, period_start
from ..images import reencode_photo
from ..live import publish_fault

router = APIRouter(prefix="/api/kvarovi", tags=["kvarovi"])
logger = logging.getLogger("adixpaper")

OPEN_STATUSES = {StatusKvaraEnum.novo, StatusKvaraEnum.u_rjesavanju, StatusKvaraEnum.na_cekanju}


def _enum_value(value) -> str:
    raw = getattr(value, "value", value)
    return str(raw) if raw is not None else ""


OPEN_FILTER = "otvoreni"
MAX_PRILOGA = 6
CONFLICT_DETAIL = "Kvar je u međuvremenu izmijenjen. Učitajte najnovije stanje pa ponovite promjenu."


def _enum_param(enum_cls, value: Optional[str]):
    if not value:
        return None
    normalized = value.strip().lower()
    for item in enum_cls:
        if item.name.lower() == normalized or item.value.lower() == normalized:
            return item
    raise HTTPException(status_code=422, detail="Neispravna vrijednost filtra")


def _user_label(user: Optional[Korisnik]) -> str:
    if not user:
        return ""
    return f"{user.ime} {user.prezime}".strip() or user.username


def _can_manage(user: Korisnik) -> bool:
    return role_value(user.rola) in ("admin", "voditelj")


def _ensure_kvar_access(kvar: Kvar, user: Korisnik):
    if role_value(user.rola) == "radnik" and kvar.prijavio_id != user.id:
        raise HTTPException(status_code=404, detail="Kvar nije pronađen")


def _kvar_for(kvar: Kvar, user: Korisnik) -> KvarOut:
    """The fault as this user may see it: workers never get internal notes, costs or parts."""
    out = KvarOut.model_validate(kvar)
    if not _can_manage(user):
        out = out.model_copy(update=KVAR_INTERNAL_FIELDS)
    return out


def _sla_ratio(kvar: Kvar, now: Optional[datetime] = None) -> Optional[float]:
    if kvar.status == StatusKvaraEnum.rijeseno:
        return None
    limit = SLA_SATI.get(_enum_value(kvar.prioritet))
    if not limit or not kvar.datum_prijave:
        return None
    now = now or utcnow()
    elapsed = (now - as_utc(kvar.datum_prijave)).total_seconds() / 3600
    return elapsed / limit


def _sla_state(kvar: Kvar, now: Optional[datetime] = None) -> str:
    ratio = _sla_ratio(kvar, now)
    if ratio is None:
        return "nema"
    if ratio >= 1:
        return "kasni"
    if ratio >= 0.75:
        return "uskoro"
    return "ok"


SLA_LABELS = {"kasni": "Kasni", "uskoro": "Ističe uskoro", "ok": "U roku", "nema": "—"}


def _sla_label(kvar: Kvar) -> str:
    return SLA_LABELS[_sla_state(kvar)]


def _add_event(
    db: Session,
    kvar: Kvar,
    user: Optional[Korisnik],
    tip: str,
    naslov: str,
    opis: Optional[str] = None,
    podaci: Optional[dict] = None,
    interno: bool = False,
) -> KvarDogadjaj:
    event = KvarDogadjaj(
        kvar_id=kvar.id,
        korisnik_id=user.id if user else None,
        tip=tip,
        naslov=naslov,
        opis=opis,
        podaci=podaci,
        interno=interno,
    )
    db.add(event)
    return event


def _query_kvarovi(
    db: Session,
    current_user: Korisnik,
    status: Optional[str] = None,
    prioritet: Optional[str] = None,
    oprema_id: Optional[int] = None,
    sla: Optional[str] = None,
    with_parts: bool = True,
    cekanje: Optional[str] = None,
    preuzeo: Optional[str] = None,
    q: Optional[str] = None,
    sort: Optional[str] = None,
    limit: Optional[int] = None,
    offset: int = 0,
    rijeseno: Optional[str] = None,
):
    """Faults matching the filters, as (page, total). Without `limit`, the page is every match."""
    query = db.query(Kvar).options(
        joinedload(Kvar.oprema),
        joinedload(Kvar.komponenta),
        joinedload(Kvar.prijavio),
        joinedload(Kvar.voditelj),
    )
    if with_parts:
        query = query.options(joinedload(Kvar.utroseni_dijelovi).joinedload(KvarRezervniDio.rezervni_dio).joinedload(RezervniDio.dobavljac))
    if role_value(current_user.rola) == "radnik":
        query = query.filter(Kvar.prijavio_id == current_user.id)
    if status and status.strip().lower() == OPEN_FILTER:
        query = query.filter(Kvar.status.in_(list(OPEN_STATUSES)))
    elif status:
        query = query.filter(Kvar.status == _enum_param(StatusKvaraEnum, status))
    if prioritet:
        query = query.filter(Kvar.prioritet == _enum_param(PrioritetEnum, prioritet))
    if oprema_id:
        query = query.filter(Kvar.oprema_id == oprema_id)
    if cekanje:
        if cekanje not in RAZLOZI_CEKANJA:
            raise HTTPException(status_code=422, detail="Neispravna vrijednost filtra")
        query = query.filter(Kvar.status == StatusKvaraEnum.na_cekanju, Kvar.razlog_cekanja == cekanje)
    if rijeseno:
        # Resolved within an overview period ("7d", "mjesec", …), so the overview's count links here.
        since = period_start(rijeseno)
        if since is None:
            raise HTTPException(status_code=422, detail="Neispravna vrijednost filtra")
        query = query.filter(Kvar.status == StatusKvaraEnum.rijeseno, Kvar.datum_rjesavanja >= since)
    if preuzeo:
        if preuzeo == "ja":
            query = query.filter(Kvar.voditelj_id == current_user.id)
        elif preuzeo == "nitko":
            query = query.filter(Kvar.voditelj_id.is_(None))
        elif preuzeo.isdigit():
            query = query.filter(Kvar.voditelj_id == int(preuzeo))
        else:
            raise HTTPException(status_code=422, detail="Neispravna vrijednost filtra")
    text_query = (q or "").strip().lstrip("#")
    if text_query:
        if text_query.isdigit():
            query = query.filter(Kvar.id == int(text_query))
        else:
            pattern = f"%{text_query.lower()}%"
            query = query.outerjoin(Oprema, Kvar.oprema_id == Oprema.id).outerjoin(
                KomponentaOpreme, Kvar.komponenta_id == KomponentaOpreme.id
            ).filter(or_(
                func.lower(Kvar.naslov).like(pattern),
                func.lower(Oprema.naziv).like(pattern),
                func.lower(KomponentaOpreme.naziv).like(pattern),
            ))
    if sort not in (None, "", "novo", "staro", "rizik", "rijeseno"):
        raise HTTPException(status_code=422, detail="Neispravna vrijednost filtra")
    if sort == "rijeseno":
        # Most recently resolved first (the resolved-faults view); unresolved ones last.
        query = query.order_by(Kvar.datum_rjesavanja.is_(None), desc(Kvar.datum_rjesavanja), desc(Kvar.id))
    elif sort == "staro":
        query = query.order_by(Kvar.datum_prijave.asc(), Kvar.id.asc())
    else:
        query = query.order_by(desc(Kvar.datum_prijave), desc(Kvar.id))

    wanted_sla = (sla or "").strip().lower()
    if wanted_sla and wanted_sla not in SLA_LABELS:
        raise HTTPException(status_code=422, detail="Neispravna vrijednost filtra")
    if wanted_sla or sort == "rizik":
        # Time-left filters and ordering depend on "now": computed here, then paged.
        records = query.all()
        if wanted_sla:
            records = [k for k in records if _sla_state(k) == wanted_sla]
        if sort == "rizik":
            far = datetime.max.replace(tzinfo=timezone.utc)
            records.sort(key=lambda k: (k.rok if k.otvoren and k.rok else far))
        total = len(records)
        return (records[offset:offset + limit] if limit else records), total
    if limit:
        total = query.order_by(None).with_entities(func.count(func.distinct(Kvar.id))).scalar() or 0
        return query.offset(offset).limit(limit).all(), total
    records = query.all()
    return records, len(records)


PUSH_TTL_SECONDS = 24 * 3600  # a phone that is off or out of signal still gets it within a day


def _send_push_notifications(title: str, body: str, url: str = "/", user_ids: Optional[List[int]] = None, tag: Optional[str] = None) -> dict:
    """Runs as a background task, after the request's session is closed, so it opens its own.
    Without `user_ids`, goes to every active lead and admin (new faults). The payload holds only
    a short title and line plus an in-app link: the app shows details after sign-in. `tag` lets
    the phone replace an older notification about the same thing instead of stacking a copy.
    Returns {"devices", "sent", "expired", "failed"} and logs the same, without endpoints or keys."""
    from ..config import settings

    result = {"devices": 0, "sent": 0, "expired": 0, "failed": 0}
    if not settings.VAPID_PRIVATE_KEY:
        return result

    db = SessionLocal()
    try:
        from pywebpush import webpush, WebPushException

        query = db.query(PushPretplata).join(Korisnik, PushPretplata.korisnik_id == Korisnik.id).filter(Korisnik.aktivan == True)
        if user_ids is None:
            query = query.filter(Korisnik.rola.in_([RolaEnum.admin, RolaEnum.voditelj]))
        else:
            query = query.filter(Korisnik.id.in_(user_ids))
        pretplate = query.all()
        result["devices"] = len(pretplate)
        mrtve = []
        for p in pretplate:
            service = urlsplit(p.endpoint).hostname or "?"
            try:
                webpush(
                    subscription_info={
                        "endpoint": p.endpoint,
                        "keys": {"p256dh": p.p256dh, "auth": p.auth}
                    },
                    data=json.dumps({"title": title, "body": body, "url": url, "tag": tag or url}),
                    vapid_private_key=settings.VAPID_PRIVATE_KEY,
                    vapid_claims={"sub": settings.VAPID_EMAIL},
                    # pywebpush defaults to TTL 0 ("now or never"): a locked or offline phone would miss it.
                    ttl=PUSH_TTL_SECONDS,
                    headers={"Urgency": "high"},
                )
                result["sent"] += 1
            except WebPushException as e:
                # A requests.Response is falsy for 4xx/5xx, so compare with None, not truthiness.
                status = e.response.status_code if e.response is not None else None
                if status in (404, 410):
                    mrtve.append(p.id)
                    result["expired"] += 1
                else:
                    result["failed"] += 1
                    reason = (e.response.text[:200] if e.response is not None else str(e))
                    logger.warning("Push to %s failed (user %s): HTTP %s %s", service, p.korisnik_id, status, reason)
        for pid in mrtve:
            db.query(PushPretplata).filter(PushPretplata.id == pid).delete()
        if mrtve:
            db.commit()
        logger.info("Push %r: %s device(s), %s sent, %s expired, %s failed", url, result["devices"], result["sent"], result["expired"], result["failed"])
    except Exception:
        logger.exception("Sending push notifications failed")
    finally:
        db.close()
    return result


def _notify_reporter(background_tasks: BackgroundTasks, kvar: Kvar, actor: Korisnik, body: str):
    """Tell the worker who reported the fault that something they can see changed."""
    if kvar.prijavio_id and kvar.prijavio_id != actor.id:
        background_tasks.add_task(
            _send_push_notifications, f"Kvar #{kvar.id}: {kvar.naslov}"[:120], body, f"/radnik/kvar/{kvar.id}", [kvar.prijavio_id]
        )


def _normalized(value):
    """Treat None/""/0 alike so a re-saved empty field is not logged as a change."""
    if value in (None, ""):
        return None
    if isinstance(value, (int, float, Decimal)):
        return None if Decimal(str(value)) == 0 else Decimal(str(value)).normalize()
    return value


def _parts_signature(items) -> list:
    rows = []
    for item in items:
        get = item.get if isinstance(item, dict) else lambda key, _i=item: getattr(_i, key)
        rows.append((int(get("rezervni_dio_id")), Decimal(str(get("kolicina"))).normalize(), get("napomena") or None))
    return sorted(rows, key=lambda row: (row[0], row[1], row[2] or ""))


def _load_kvar(kvar_id: int, db: Session) -> Kvar:
    kvar = db.query(Kvar).options(
        joinedload(Kvar.oprema),
        joinedload(Kvar.komponenta),
        joinedload(Kvar.prijavio),
        joinedload(Kvar.voditelj),
        joinedload(Kvar.utroseni_dijelovi).joinedload(KvarRezervniDio.rezervni_dio).joinedload(RezervniDio.dobavljac),
    ).filter(Kvar.id == kvar_id).first()
    if not kvar:
        raise HTTPException(status_code=404, detail="Kvar nije pronađen")
    return kvar


@router.get("", response_model=List[KvarListOut])
def list_kvarovi(
    response: Response,
    status: Optional[str] = None,
    prioritet: Optional[str] = None,
    oprema_id: Optional[int] = None,
    sla: Optional[str] = None,
    cekanje: Optional[str] = None,
    preuzeo: Optional[str] = Query(None, max_length=20),
    q: Optional[str] = Query(None, max_length=200),
    sort: Optional[str] = None,
    limit: Optional[int] = Query(None, ge=1, le=200),
    offset: int = Query(0, ge=0),
    rijeseno: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(get_current_user)
):
    """Faults the user may see. With `limit`, one page; the total is in the X-Total-Count header."""
    records, total = _query_kvarovi(
        db, current_user, status, prioritet, oprema_id, sla, with_parts=False,
        cekanje=cekanje, preuzeo=preuzeo, q=q, sort=sort, limit=limit, offset=offset, rijeseno=rijeseno,
    )
    response.headers["X-Total-Count"] = str(total)
    return records


@router.get("/novi", response_model=List[KvarNovoOut])
def list_new_kvarovi(
    since_id: int = Query(0, ge=0),
    db: Session = Depends(get_db),
    _: Korisnik = Depends(require_voditelj_or_admin),
):
    """Lightweight feed polled by the admin shell to announce newly reported faults."""
    return (
        db.query(Kvar)
        .options(joinedload(Kvar.oprema))
        .filter(Kvar.id > since_id)
        .order_by(Kvar.id.desc())
        .limit(20)
        .all()
    )


@router.get("/na-masini/{oprema_id}", response_model=List[KvarNaMasiniOut])
def list_open_on_machine(
    oprema_id: int,
    db: Session = Depends(get_db),
    _: Korisnik = Depends(get_current_user),
):
    """Open faults on a machine, so someone reporting can see it may already be known.

    Everyone may call it, workers included: it returns only number, title, status, component
    and time, never who reported or handles it, costs or notes.
    """
    return (
        db.query(Kvar)
        .options(joinedload(Kvar.komponenta))
        .filter(Kvar.oprema_id == oprema_id, Kvar.status.in_(list(OPEN_STATUSES)))
        .order_by(desc(Kvar.datum_prijave))
        .limit(10)
        .all()
    )


@router.get("/export.xlsx")
def export_kvarovi_xlsx(
    status: Optional[str] = None,
    prioritet: Optional[str] = None,
    oprema_id: Optional[int] = None,
    sla: Optional[str] = None,
    cekanje: Optional[str] = None,
    preuzeo: Optional[str] = Query(None, max_length=20),
    q: Optional[str] = Query(None, max_length=200),
    sort: Optional[str] = None,
    rijeseno: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(require_voditelj_or_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    records, _ = _query_kvarovi(db, current_user, status, prioritet, oprema_id, sla, cekanje=cekanje, preuzeo=preuzeo, q=q, sort=sort, rijeseno=rijeseno)
    wb = Workbook()
    ws = wb.active
    ws.title = "Kvarovi"
    headers = [
        "ID", "Naslov", "Oprema", "Lokacija", "Prioritet", "Status", "Rok",
        "Prijavio", "Voditelj", "Datum prijave", "Datum pocetka", "Datum rjesavanja",
        "Trajanje min", "Troskovi EUR", "Materijal", "Sto je uradjeno"
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")

    for k in records:
        dijelovi = "; ".join(
            f"{u.rezervni_dio.naziv} x {float(u.kolicina or 0):g}"
            for u in (k.utroseni_dijelovi or [])
            if u.rezervni_dio
        )
        ws.append([
            k.id,
            k.naslov,
            k.oprema.naziv if k.oprema else "",
            k.oprema.lokacija if k.oprema else "",
            _enum_value(k.prioritet),
            _enum_value(k.status),
            _sla_label(k),
            _user_label(k.prijavio),
            _user_label(k.voditelj),
            local_naive(k.datum_prijave),
            local_naive(k.datum_pocetka),
            local_naive(k.datum_rjesavanja),
            k.trajanje_min or 0,
            float(k.troskovi or 0),
            dijelovi or k.materijal or "",
            k.sto_je_uradjeno or "",
        ])

    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 10), 42)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    headers = {"Content-Disposition": 'attachment; filename="adix-paper-kvarovi.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.get("/export-dijelovi.xlsx")
def export_utroseni_dijelovi_xlsx(
    status: Optional[str] = None,
    prioritet: Optional[str] = None,
    oprema_id: Optional[int] = None,
    sla: Optional[str] = None,
    cekanje: Optional[str] = None,
    preuzeo: Optional[str] = Query(None, max_length=20),
    q: Optional[str] = Query(None, max_length=200),
    sort: Optional[str] = None,
    rijeseno: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(require_voditelj_or_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    records, _ = _query_kvarovi(db, current_user, status, prioritet, oprema_id, sla, cekanje=cekanje, preuzeo=preuzeo, q=q, sort=sort, rijeseno=rijeseno)
    kvar_ids = [k.id for k in records]
    utrosci = []
    if kvar_ids:
        utrosci = db.query(KvarRezervniDio).options(
            joinedload(KvarRezervniDio.kvar).joinedload(Kvar.oprema),
            joinedload(KvarRezervniDio.kvar).joinedload(Kvar.komponenta),
            joinedload(KvarRezervniDio.rezervni_dio).joinedload(RezervniDio.dobavljac),
        ).filter(KvarRezervniDio.kvar_id.in_(kvar_ids)).all()

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    grouped = {}
    for item in utrosci:
        supplier = item.rezervni_dio.dobavljac.naziv if item.rezervni_dio and item.rezervni_dio.dobavljac else "Bez dobavljaca"
        grouped.setdefault(supplier, []).append(item)

    if not grouped:
        grouped["Nema utroska"] = []

    for supplier, rows in grouped.items():
        title = supplier[:31] or "Dobavljac"
        ws = wb.create_sheet(title=title)
        headers = [
            "Dobavljac", "Kvar ID", "Datum", "Oprema", "Komponenta", "Dio",
            "Sifra", "Kolicina", "Jedinica", "Zaliha", "Cijena EUR", "Ukupno EUR", "Napomena"
        ]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="111827")

        for item in rows:
            dio = item.rezervni_dio
            kvar = item.kvar
            qty = float(item.kolicina or 0)
            price = float(dio.cijena or 0) if dio else 0
            ws.append([
                supplier,
                kvar.id if kvar else "",
                local_naive(kvar.datum_prijave) if kvar else "",
                kvar.oprema.naziv if kvar and kvar.oprema else "",
                kvar.komponenta.naziv if kvar and kvar.komponenta else "",
                dio.naziv if dio else "",
                dio.sifra if dio else "",
                qty,
                dio.jedinica if dio else "",
                float(dio.zaliha or 0) if dio and dio.zaliha is not None else "",
                price,
                qty * price,
                item.napomena or "",
            ])

        for column in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in column)
            ws.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 10), 42)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    headers = {"Content-Disposition": 'attachment; filename="adix-paper-utroseni-dijelovi.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.get("/export.pdf")
def export_kvarovi_pdf(
    status: Optional[str] = None,
    prioritet: Optional[str] = None,
    oprema_id: Optional[int] = None,
    sla: Optional[str] = None,
    cekanje: Optional[str] = None,
    preuzeo: Optional[str] = Query(None, max_length=20),
    q: Optional[str] = Query(None, max_length=200),
    sort: Optional[str] = None,
    rijeseno: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(require_voditelj_or_admin),
):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    from ..pdf import BRAND_BLUE, register_fonts, stylesheet

    records, _ = _query_kvarovi(db, current_user, status, prioritet, oprema_id, sla, cekanje=cekanje, preuzeo=preuzeo, q=q, sort=sort, rijeseno=rijeseno)
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    regular, bold = register_fonts()
    styles = stylesheet()
    story = [
        Paragraph("Adix Paper – izvještaj kvarova", styles["Title"]),
        Paragraph(local_naive(utcnow()).strftime("Generirano: %d.%m.%Y. %H:%M"), styles["Normal"]),
        Spacer(1, 12),
    ]

    rows = [["ID", "Naslov", "Mašina", "Prioritet", "Status", "Rok", "Prijava", "Trošak €"]]
    for k in records:
        rows.append([
            str(k.id),
            Paragraph((k.naslov or "")[:48], styles["BodyText"]),
            Paragraph((k.oprema.naziv if k.oprema else "")[:36], styles["BodyText"]),
            _enum_value(k.prioritet),
            _enum_value(k.status),
            _sla_label(k),
            local_naive(k.datum_prijave).strftime("%d.%m.%Y") if k.datum_prijave else "",
            f"{float(k.troskovi or 0):.2f}",
        ])

    table = Table(rows, repeatRows=1, colWidths=[34, 180, 150, 72, 86, 58, 72, 58])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND_BLUE)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("FONTNAME", (0, 0), (-1, -1), regular),
        ("FONTNAME", (0, 0), (-1, 0), bold),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(table)
    doc.build(story)
    output.seek(0)
    headers = {"Content-Disposition": 'attachment; filename="adix-paper-kvarovi.pdf"'}
    return StreamingResponse(output, media_type="application/pdf", headers=headers)


@router.get("/{kvar_id}", response_model=KvarOut)
def get_kvar(kvar_id: int, db: Session = Depends(get_db), current_user: Korisnik = Depends(get_current_user)):
    kvar = _load_kvar(kvar_id, db)
    _ensure_kvar_access(kvar, current_user)
    return _kvar_for(kvar, current_user)


def _fault_for_key(db: Session, key: Optional[str], user: Korisnik) -> Optional[Kvar]:
    """The fault this user already created with this report key; 409 if another user's."""
    if not key:
        return None
    existing = db.query(Kvar).filter(Kvar.klijent_kljuc == key).first()
    if existing is not None and existing.prijavio_id != user.id:
        raise HTTPException(status_code=409, detail="Ključ prijave već je iskorišten")
    return existing


@router.post("", response_model=KvarOut, status_code=201)
def create_kvar(
    data: KvarCreate,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(get_current_user)
):
    # Workers (and maintenance leads, on the floor) report faults; administrators manage them.
    if role_value(current_user.rola) == "admin":
        raise HTTPException(status_code=403, detail="Administrator ne prijavljuje kvarove")
    oprema = db.query(Oprema).filter(Oprema.id == data.oprema_id).first()
    if not oprema:
        raise HTTPException(status_code=404, detail="Oprema nije pronađena")
    if data.komponenta_id:
        komponenta = db.query(KomponentaOpreme).filter(
            KomponentaOpreme.id == data.komponenta_id,
            KomponentaOpreme.oprema_id == data.oprema_id,
            KomponentaOpreme.aktivna == True,
        ).first()
        if not komponenta:
            raise HTTPException(status_code=404, detail="Komponenta nije pronađena za odabranu opremu")

    if data.klijent_kljuc:
        existing = _fault_for_key(db, data.klijent_kljuc, current_user)
        if existing is not None:
            return _kvar_for(_load_kvar(existing.id, db), current_user)

    kvar = Kvar(
        oprema_id=data.oprema_id,
        komponenta_id=data.komponenta_id,
        prijavio_id=current_user.id,
        naslov=data.naslov,
        opis=data.opis,
        prioritet=data.prioritet,
        status=StatusKvaraEnum.novo,
        klijent_kljuc=data.klijent_kljuc,
    )
    db.add(kvar)
    try:
        db.flush()
    except IntegrityError:
        # The same report arrived twice at once: the other request won the unique key.
        db.rollback()
        existing = _fault_for_key(db, data.klijent_kljuc, current_user)
        if existing is None:
            raise
        return _kvar_for(_load_kvar(existing.id, db), current_user)
    _add_event(
        db,
        kvar,
        current_user,
        "created",
        "Kvar prijavljen",
        data.opis,
        {"prioritet": data.prioritet.value, "status": StatusKvaraEnum.novo.value, "komponenta_id": data.komponenta_id},
    )
    db.commit()
    db.refresh(kvar)
    publish_fault(kvar, "created")

    background_tasks.add_task(
        _send_push_notifications,
        f"Novi kvar [{data.prioritet.value}]",
        f"{oprema.naziv}: {data.naslov}",
        f"/admin/kvarovi/{kvar.id}",
        None,
        f"kvar-{kvar.id}-novi",
    )

    return _kvar_for(_load_kvar(kvar.id, db), current_user)


@router.get("/{kvar_id}/dogadjaji", response_model=List[KvarDogadjajOut])
def list_kvar_dogadjaji(
    kvar_id: int,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(get_current_user),
):
    kvar = _load_kvar(kvar_id, db)
    _ensure_kvar_access(kvar, current_user)

    q = db.query(KvarDogadjaj).options(joinedload(KvarDogadjaj.korisnik)).filter(KvarDogadjaj.kvar_id == kvar_id)
    if not _can_manage(current_user):
        q = q.filter(KvarDogadjaj.interno == False)
    return q.order_by(KvarDogadjaj.kreiran.asc(), KvarDogadjaj.id.asc()).all()


@router.post("/{kvar_id}/komentari", response_model=KvarDogadjajOut, status_code=201)
def add_kvar_komentar(
    kvar_id: int,
    data: KvarKomentarCreate,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(get_current_user),
):
    tekst = data.tekst.strip()
    if not tekst:
        raise HTTPException(status_code=400, detail="Komentar ne može biti prazan")

    kvar = _load_kvar(kvar_id, db)
    _ensure_kvar_access(kvar, current_user)
    interno = bool(data.interno and _can_manage(current_user))
    event = _add_event(
        db,
        kvar,
        current_user,
        "komentar",
        "Interni komentar" if interno else "Komentar",
        tekst,
        interno=interno,
    )
    if not interno:
        kvar.azurirano = utcnow()
    db.commit()
    db.refresh(event)
    publish_fault(kvar, "comment", internal=interno)
    if not interno:
        _notify_reporter(background_tasks, kvar, current_user, f"Komentar: {tekst[:140]}")
    return event


@router.put("/{kvar_id}", response_model=KvarOut)
def update_kvar(
    kvar_id: int,
    data: KvarVoditeljUpdate,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(require_voditelj_or_admin)
):
    # Row lock (PostgreSQL): two saves of the same fault are serialised, so the version check
    # below cannot pass for both of them.
    kvar = db.query(Kvar).filter(Kvar.id == kvar_id).with_for_update().first()
    if not kvar:
        raise HTTPException(status_code=404, detail="Kvar nije pronađen")

    update_data = data.model_dump(exclude_unset=True)
    expected_version = update_data.pop("verzija", None)
    if expected_version is not None and expected_version != kvar.verzija:
        raise HTTPException(status_code=409, detail=CONFLICT_DETAIL)
    for required in ("status", "prioritet"):
        if update_data.get(required, "") is None:
            update_data.pop(required)
    parts_payload = update_data.pop("utroseni_dijelovi", None)
    waiting_reason = update_data.pop("razlog_cekanja", None)
    old_waiting_reason = kvar.razlog_cekanja
    old_waiting_since = kvar.cekanje_od
    old_status = _enum_value(kvar.status)
    old_prioritet = _enum_value(kvar.prioritet)
    old_voditelj_id = kvar.voditelj_id
    old_komponenta_id = kvar.komponenta_id
    detail_fields = ("sto_je_uradjeno", "materijal", "troskovi", "trajanje_min", "napomene_voditelja")
    old_details = {field: getattr(kvar, field) for field in detail_fields}

    if "komponenta_id" in update_data and update_data["komponenta_id"]:
        component = db.query(KomponentaOpreme).filter(
            KomponentaOpreme.id == update_data["komponenta_id"],
            KomponentaOpreme.oprema_id == kvar.oprema_id,
        ).first()
        if not component:
            raise HTTPException(status_code=404, detail="Komponenta nije pronađena za odabranu opremu")

    new_assignee = update_data.get("voditelj_id")
    if new_assignee is not None and new_assignee != kvar.voditelj_id:
        assignee = db.query(Korisnik).filter(Korisnik.id == new_assignee).first()
        if not assignee or not assignee.aktivan or role_value(assignee.rola) not in ("admin", "voditelj"):
            raise HTTPException(status_code=400, detail="Kvar se može dodijeliti samo aktivnom voditelju ili administratoru")

    for field, value in update_data.items():
        setattr(kvar, field, value)

    if "status" in update_data and _enum_value(update_data["status"]) != old_status:
        new_status = update_data["status"]
        if new_status == StatusKvaraEnum.u_rjesavanju:
            if not kvar.datum_pocetka:
                kvar.datum_pocetka = utcnow()
            # Taking a fault makes you its owner unless someone else already is. Applied after
            # the payload so a form that still says "unassigned" cannot undo it.
            if not kvar.voditelj_id:
                kvar.voditelj_id = current_user.id
        if new_status == StatusKvaraEnum.rijeseno:
            if not (kvar.sto_je_uradjeno or "").strip():
                raise HTTPException(status_code=400, detail="Prije zatvaranja kvara upišite što je napravljeno")
            if not kvar.datum_rjesavanja:
                kvar.datum_rjesavanja = utcnow()
        else:
            kvar.datum_rjesavanja = None

    # On hold always carries a reason and a start time; leaving it clears both.
    if kvar.status == StatusKvaraEnum.na_cekanju:
        if old_status != StatusKvaraEnum.na_cekanju.value:
            if not waiting_reason:
                raise HTTPException(status_code=400, detail="Odaberite razlog čekanja")
            kvar.razlog_cekanja = waiting_reason
            kvar.cekanje_od = utcnow()
        elif waiting_reason and waiting_reason != kvar.razlog_cekanja:
            kvar.razlog_cekanja = waiting_reason
    else:
        kvar.razlog_cekanja = None
        kvar.cekanje_od = None

    db.flush()

    new_status = _enum_value(kvar.status)
    new_prioritet = _enum_value(kvar.prioritet)

    if old_status != new_status:
        podaci = {"old": old_status, "new": new_status}
        if kvar.razlog_cekanja:
            podaci["razlog"] = kvar.razlog_cekanja
        if old_waiting_since is not None and kvar.cekanje_od is None:
            # How long it was on hold, and why: the source for "waiting time by reason" in Analytics.
            podaci["cekanje_razlog"] = old_waiting_reason
            podaci["cekanje_min"] = round((utcnow() - as_utc(old_waiting_since)).total_seconds() / 60)
        _add_event(
            db,
            kvar,
            current_user,
            "status",
            f"Status promijenjen: {old_status} -> {new_status}",
            podaci=podaci,
        )
    elif old_waiting_reason and kvar.razlog_cekanja and old_waiting_reason != kvar.razlog_cekanja:
        _add_event(
            db,
            kvar,
            current_user,
            "waiting",
            "Promijenjen razlog čekanja",
            podaci={"old": old_waiting_reason, "new": kvar.razlog_cekanja},
        )
    if old_prioritet != new_prioritet:
        _add_event(
            db,
            kvar,
            current_user,
            "prioritet",
            f"Prioritet promijenjen: {old_prioritet} -> {new_prioritet}",
            podaci={"old": old_prioritet, "new": new_prioritet},
        )
    if old_voditelj_id != kvar.voditelj_id:
        _add_event(
            db,
            kvar,
            current_user,
            "assignment",
            "Promijenjen voditelj kvara",
            podaci={"old": old_voditelj_id, "new": kvar.voditelj_id},
        )
    if old_komponenta_id != kvar.komponenta_id:
        _add_event(
            db,
            kvar,
            current_user,
            "component",
            "Promijenjena komponenta kvara",
            podaci={"old": old_komponenta_id, "new": kvar.komponenta_id},
        )

    changed_details = sorted(
        field for field in detail_fields
        if field in update_data and _normalized(old_details[field]) != _normalized(getattr(kvar, field))
    )
    if changed_details:
        _add_event(
            db,
            kvar,
            current_user,
            "details",
            "Ažurirani detalji rješavanja",
            podaci={"fields": changed_details},
            interno=changed_details == ["napomene_voditelja"],
        )

    if parts_payload is not None and _parts_signature(parts_payload) == _parts_signature(
        [{"rezervni_dio_id": u.rezervni_dio_id, "kolicina": u.kolicina, "napomena": u.napomena} for u in kvar.utroseni_dijelovi]
    ):
        parts_payload = None  # unchanged: keep rows and avoid a spurious timeline entry

    if parts_payload is not None:
        db.query(KvarRezervniDio).filter(KvarRezervniDio.kvar_id == kvar.id).delete()
        for item in parts_payload:
            rezervni_dio_id = item.get("rezervni_dio_id") if isinstance(item, dict) else item.rezervni_dio_id
            kolicina = item.get("kolicina") if isinstance(item, dict) else item.kolicina
            napomena = item.get("napomena") if isinstance(item, dict) else item.napomena

            if kolicina <= 0:
                raise HTTPException(status_code=400, detail="Količina rezervnog dijela mora biti veća od nule")
            if not db.query(RezervniDio.id).filter(RezervniDio.id == rezervni_dio_id, RezervniDio.aktivan == True).first():
                raise HTTPException(status_code=404, detail="Rezervni dio nije pronađen")
            db.add(KvarRezervniDio(
                kvar_id=kvar.id,
                rezervni_dio_id=rezervni_dio_id,
                kolicina=kolicina,
                napomena=napomena,
            ))
        _add_event(
            db,
            kvar,
            current_user,
            "parts",
            "Ažurirani utrošeni rezervni dijelovi",
            podaci={"count": len(parts_payload)},
        )

    kvar.verzija = (kvar.verzija or 1) + 1
    visible_change = old_status != new_status or old_voditelj_id != kvar.voditelj_id or kvar.razlog_cekanja != old_waiting_reason
    if visible_change:
        kvar.azurirano = utcnow()
    db.commit()
    # Workers hear only about changes they can see (status, handler, waiting reason).
    publish_fault(kvar, "updated", internal=not visible_change)
    if old_status != new_status:
        _notify_reporter(background_tasks, kvar, current_user, f"Status: {new_status}")
    elif old_voditelj_id != kvar.voditelj_id and kvar.voditelj:
        _notify_reporter(background_tasks, kvar, current_user, f"Preuzeo: {_user_label(kvar.voditelj)}")
    return _load_kvar(kvar_id, db)


@router.delete("/{kvar_id}", status_code=204)
def delete_kvar(
    kvar_id: int,
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(require_admin),
):
    """Removes a fault and its history. Admin only (e.g. a test or duplicate entry); not offered in the UI."""
    kvar = db.query(Kvar).filter(Kvar.id == kvar_id).first()
    if not kvar:
        raise HTTPException(status_code=404, detail="Kvar nije pronađen")
    reporter_id = kvar.prijavio_id
    db.delete(kvar)
    db.commit()
    publish_fault(SimpleNamespace(id=kvar_id, prijavio_id=reporter_id), "deleted")
    logger.warning("Fault %s deleted by %s", kvar_id, current_user.username)


@router.get("/{kvar_id}/prilozi", response_model=List[KvarPrilogOut])
def list_prilozi(kvar_id: int, db: Session = Depends(get_db), current_user: Korisnik = Depends(get_current_user)):
    kvar = _load_kvar(kvar_id, db)
    _ensure_kvar_access(kvar, current_user)
    return db.query(KvarPrilog).options(joinedload(KvarPrilog.korisnik)).filter(KvarPrilog.kvar_id == kvar_id).order_by(KvarPrilog.id).all()


@router.post("/{kvar_id}/prilozi", response_model=KvarPrilogOut, status_code=201)
def add_prilog(
    kvar_id: int,
    datoteka: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: Korisnik = Depends(get_current_user),
):
    """Attach a photo. The reporter (worker) and leads/admins may add up to MAX_PRILOGA per fault."""
    kvar = _load_kvar(kvar_id, db)
    _ensure_kvar_access(kvar, current_user)
    if db.query(func.count(KvarPrilog.id)).filter(KvarPrilog.kvar_id == kvar_id).scalar() >= MAX_PRILOGA:
        raise HTTPException(status_code=400, detail=f"Kvar može imati najviše {MAX_PRILOGA} fotografija")
    raw = datoteka.file.read(12 * 1024 * 1024 + 1)
    jpeg, width, height = reencode_photo(raw)
    name = (datoteka.filename or "fotografija").rsplit("/", 1)[-1].rsplit("\\", 1)[-1][:150] or "fotografija"
    prilog = KvarPrilog(
        kvar_id=kvar_id,
        korisnik_id=current_user.id,
        naziv=name,
        mime="image/jpeg",
        velicina=len(jpeg),
        sirina=width,
        visina=height,
        sadrzaj=jpeg,
    )
    db.add(prilog)
    db.flush()
    _add_event(db, kvar, current_user, "prilog", "Dodana fotografija", podaci={"prilog_id": prilog.id})
    kvar.azurirano = utcnow()
    db.commit()
    db.refresh(prilog)
    publish_fault(kvar, "photo")
    return prilog


@router.get("/{kvar_id}/prilozi/{prilog_id}")
def get_prilog(kvar_id: int, prilog_id: int, db: Session = Depends(get_db), current_user: Korisnik = Depends(get_current_user)):
    kvar = _load_kvar(kvar_id, db)
    _ensure_kvar_access(kvar, current_user)
    prilog = db.query(KvarPrilog).filter(KvarPrilog.id == prilog_id, KvarPrilog.kvar_id == kvar_id).first()
    if not prilog:
        raise HTTPException(status_code=404, detail="Fotografija nije pronađena")
    return Response(
        content=prilog.sadrzaj,
        media_type=prilog.mime,
        headers={"Cache-Control": "private, max-age=86400", "Content-Disposition": "inline"},
    )


@router.delete("/{kvar_id}/prilozi/{prilog_id}", status_code=204)
def delete_prilog(kvar_id: int, prilog_id: int, db: Session = Depends(get_db), current_user: Korisnik = Depends(get_current_user)):
    """Leads/admins may remove any photo; a worker only their own, while the fault is open."""
    kvar = _load_kvar(kvar_id, db)
    _ensure_kvar_access(kvar, current_user)
    prilog = db.query(KvarPrilog).filter(KvarPrilog.id == prilog_id, KvarPrilog.kvar_id == kvar_id).first()
    if not prilog:
        raise HTTPException(status_code=404, detail="Fotografija nije pronađena")
    if not _can_manage(current_user) and (prilog.korisnik_id != current_user.id or not kvar.otvoren):
        raise HTTPException(status_code=403, detail="Nemate ovlasti za ovu radnju")
    db.delete(prilog)
    _add_event(db, kvar, current_user, "prilog", "Uklonjena fotografija", podaci={"prilog_id": prilog_id})
    db.commit()
    publish_fault(kvar, "photo")
