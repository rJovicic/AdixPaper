"""Spare parts keep the unit price they had when they were used: editing the catalogue price
later never changes past costs (owner, 2026-10-29)."""
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from app import bootstrap
from app.database import SessionLocal
from app.routers.dashboard import _summary
from app.timeutil import utcnow
from tests.test_kvarovi import report


def _part(client, admin, price):
    response = client.post("/api/rezervni-dijelovi", json={"naziv": f"Ležaj {price}", "cijena": price}, headers=admin)
    assert response.status_code == 201, response.text
    return response.json()


def _set_price(client, admin, part, price):
    assert client.put(f"/api/rezervni-dijelovi/{part['id']}", json={"cijena": price}, headers=admin).status_code == 200


def test_changing_the_catalogue_price_does_not_change_past_costs(client, admin, make_user, machine):
    _, worker = make_user()
    kvar = report(client, worker, machine)
    part = _part(client, admin, "10.00")
    url = f"/api/kvarovi/{kvar['id']}"
    resolved = client.put(url, json={
        "status": "Riješeno", "sto_je_uradjeno": "Zamijenjen ležaj",
        "utroseni_dijelovi": [{"rezervni_dio_id": part["id"], "kolicina": "3"}],
    }, headers=admin).json()
    assert resolved["utroseni_dijelovi"][0]["jedinicna_cijena"] == "10.00"

    db = SessionLocal()
    try:
        start = utcnow().replace(hour=0, minute=0)
        before = _summary(db, start.replace(year=start.year - 1), utcnow()).trosak_dijelova
    finally:
        db.close()

    _set_price(client, admin, part, "25.00")

    # The fault, the analytics figure and the export still use 10.00.
    assert client.get(url, headers=admin).json()["utroseni_dijelovi"][0]["jedinicna_cijena"] == "10.00"
    db = SessionLocal()
    try:
        after = _summary(db, start.replace(year=start.year - 1), utcnow()).trosak_dijelova
    finally:
        db.close()
    assert after == before
    sheet_book = load_workbook(BytesIO(client.get(f"/api/kvarovi/export-dijelovi.xlsx?oprema_id={machine['id']}", headers=admin).content))
    rows = [row for sheet in sheet_book.worksheets for row in sheet.iter_rows(min_row=2, values_only=True) if row[1] == kvar["id"]]
    assert rows and rows[0][10] == 10.0 and rows[0][11] == 30.0

    # Re-saving the list (another quantity, another part) keeps 10.00 for the part already there
    # and gives a newly added part today's price.
    other = _part(client, admin, "4.00")
    saved = client.put(url, json={"utroseni_dijelovi": [
        {"rezervni_dio_id": part["id"], "kolicina": "2"}, {"rezervni_dio_id": other["id"], "kolicina": "1"},
    ]}, headers=admin).json()
    prices = {u["rezervni_dio_id"]: u["jedinicna_cijena"] for u in saved["utroseni_dijelovi"]}
    assert prices == {part["id"]: "10.00", other["id"]: "4.00"}


def test_upgrade_records_existing_usage_at_the_catalogue_price_and_first_responses(tmp_path, monkeypatch):
    """The additive migration on a database from the previous release (SQLite here; CI's upgrade
    job runs the same on PostgreSQL 18): existing rows are kept, prices are frozen and flagged,
    first responses come from the timeline."""
    engine = create_engine(f"sqlite:///{tmp_path / 'legacy.db'}")
    with engine.begin() as c:
        c.execute(text("CREATE TABLE rezervni_dijelovi (id INTEGER PRIMARY KEY, dobavljac_id INTEGER, cijena NUMERIC(10,2), zaliha NUMERIC(10,2))"))
        c.execute(text("CREATE TABLE kvarovi (id INTEGER PRIMARY KEY, status VARCHAR(20), datum_prijave DATETIME, datum_pocetka DATETIME, datum_rjesavanja DATETIME)"))
        c.execute(text("CREATE TABLE kvar_rezervni_dijelovi (id INTEGER PRIMARY KEY, kvar_id INTEGER, rezervni_dio_id INTEGER, kolicina NUMERIC(10,2))"))
        c.execute(text("CREATE TABLE kvar_dogadjaji (id INTEGER PRIMARY KEY, kvar_id INTEGER, tip VARCHAR(40), podaci JSON, kreiran DATETIME)"))
        c.execute(text("INSERT INTO rezervni_dijelovi VALUES (1, NULL, 12.50, NULL), (2, NULL, NULL, NULL)"))
        c.execute(text("""INSERT INTO kvarovi VALUES
            (1, 'u_rjesavanju', '2026-10-01 08:00:00', '2026-10-01 09:00:00', NULL),
            (2, 'rijeseno', '2026-10-02 08:00:00', NULL, '2026-10-02 10:00:00'),
            (3, 'novo', '2026-10-03 08:00:00', NULL, NULL),
            (4, 'novo', '2026-10-04 08:00:00', '2026-10-04 09:00:00', NULL)"""))
        # Fault 1 left "Novo" at 08:20 (on hold) before it was taken at 09:00; fault 4 was reopened.
        c.execute(text("""INSERT INTO kvar_dogadjaji VALUES
            (1, 1, 'status', '{"old": "Novo", "new": "Na čekanju"}', '2026-10-01 08:20:00'),
            (2, 1, 'status', '{"old": "Na čekanju", "new": "U rješavanju"}', '2026-10-01 09:00:00'),
            (3, 4, 'status', '{"old": "Novo", "new": "U rješavanju"}', '2026-10-04 09:00:00'),
            (4, 4, 'status', '{"old": "U rješavanju", "new": "Novo"}', '2026-10-04 11:00:00')"""))
        c.execute(text("INSERT INTO kvar_rezervni_dijelovi VALUES (1, 1, 1, 2), (2, 2, 2, 1)"))
    monkeypatch.setattr(bootstrap, "engine", engine)

    bootstrap.ensure_additive_schema()
    bootstrap.ensure_additive_schema()  # idempotent: a second start changes nothing

    with engine.connect() as c:
        responses = dict(c.execute(text("SELECT id, datum_odziva FROM kvarovi ORDER BY id")).all())
        usage = c.execute(text("SELECT id, jedinicna_cijena, cijena_naknadno, kolicina FROM kvar_rezervni_dijelovi ORDER BY id")).all()
        c.execute(text("UPDATE rezervni_dijelovi SET cijena = 99 WHERE id = 1"))
        frozen = c.execute(text("SELECT jedinicna_cijena FROM kvar_rezervni_dijelovi WHERE id = 1")).scalar()
    assert str(responses[1]).startswith("2026-10-01 08:20")  # the first move out of "Novo"
    assert str(responses[2]).startswith("2026-10-02 10:00")  # no timeline: its resolution
    assert responses[3] is None  # still waiting
    assert str(responses[4]).startswith("2026-10-04 09:00")  # reopened: keeps its first response
    assert [(row[0], float(row[1]) if row[1] is not None else None, bool(row[2]), float(row[3])) for row in usage] == [
        (1, 12.5, True, 2.0), (2, None, True, 1.0),
    ]
    assert float(frozen) == 12.5


def test_machine_cost_is_labour_plus_parts_at_the_recorded_price(client, admin, make_user, machine):
    _, worker = make_user()
    kvar = report(client, worker, machine)
    part = _part(client, admin, "7.50")
    assert client.put(f"/api/kvarovi/{kvar['id']}", json={
        "status": "Riješeno", "sto_je_uradjeno": "Popravljeno", "troskovi": "20",
        "utroseni_dijelovi": [{"rezervni_dio_id": part["id"], "kolicina": "2"}],
    }, headers=admin).status_code == 200
    _set_price(client, admin, part, "100.00")

    single = client.get(f"/api/oprema/{machine['id']}", headers=admin).json()
    listed = next(m for m in client.get("/api/oprema", headers=admin).json() if m["id"] == machine["id"])
    assert single["ukupni_troskovi"] == listed["ukupni_troskovi"] == 35.0  # 20 labour + 2 × 7.50
    # Analytics lists the ten machines with most faults: give this one enough to be among them in a
    # shared test database. The extra faults carry no cost.
    for _ in range(15):
        report(client, worker, machine)
    rows = client.get("/api/dashboard/analitika?period=30d", headers=admin).json()["masine"]
    assert next(m for m in rows if m["oprema_id"] == machine["id"])["trosak"] == 35.0
