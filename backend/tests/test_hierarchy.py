"""Machine → component → fault hierarchy (2026-10-30): data model, identifiers, QR labels,
reporting against a machine or one of its components, seeding and the guarded test reset."""
from io import BytesIO

import pypdfium2
import pytest
import zxingcpp
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app import reset_test_data
from app.database import Base, SessionLocal
from app.models import Korisnik, KomponentaOpreme, Kvar, KvarDogadjaj, KvarPrilog, KvarRezervniDio, Oprema, RezervniDio
from app.seed_assets import COMPONENTS, MACHINE, seed_amica_matrix
from tests.conftest import unique
from tests.test_kvarovi import report


def _line(client, admin, **extra):
    response = client.post("/api/oprema", json={"naziv": unique("Linija "), **extra}, headers=admin)
    assert response.status_code == 200, response.text
    return response.json()


def _component(client, admin, machine_id, **extra):
    payload = {"naziv": unique("Odmotač "), "model": "BS 2601", "proizvodac": "PCMC", **extra}
    response = client.post(f"/api/oprema/{machine_id}/komponente", json=payload, headers=admin)
    assert response.status_code == 201, response.text
    return response.json()


def _decode_pdf(content: bytes) -> list[str]:
    texts = []
    pdf = pypdfium2.PdfDocument(content)
    for page in pdf:
        image = page.render(scale=3).to_pil()
        texts += [result.text for result in zxingcpp.read_barcodes(image)]
    return texts


# --- Data model and API ------------------------------------------------------------------------

def test_a_machine_needs_only_a_name_and_gets_a_qr_identifier(client, admin):
    machine = _line(client, admin)
    assert machine["tip"] is None and machine["inventurni_broj"] is None and machine["serijski_broj"] is None
    assert len(machine["qr_oznaka"]) >= 12 and not machine["qr_oznaka"].isdigit()
    other = _line(client, admin)
    assert other["qr_oznaka"] != machine["qr_oznaka"]


def test_components_belong_to_one_machine_and_carry_their_identifiers(client, admin):
    machine = _line(client, admin)
    component = _component(client, admin, machine["id"], inventurni_broj=unique("INV-"), serijski_broj="SN 001")
    assert component["oprema_id"] == machine["id"] and component["model"] == "BS 2601" and component["proizvodac"] == "PCMC"
    assert component["qr_oznaka"] and component["qr_oznaka"] != machine["qr_oznaka"]

    # Identifiers can be entered later without touching the QR identifier.
    later = client.put(f"/api/oprema/komponente/{component['id']}", json={"inventurni_broj": unique("INV-")}, headers=admin).json()
    assert later["qr_oznaka"] == component["qr_oznaka"]

    detail = client.get(f"/api/oprema/komponente/{component['id']}", headers=admin).json()
    assert detail["oprema"]["id"] == machine["id"]
    listed = next(m for m in client.get("/api/oprema", headers=admin).json() if m["id"] == machine["id"])
    assert listed["broj_komponenti"] == 1


def test_identifiers_are_optional_but_unique_where_the_business_needs_it(client, admin):
    inventory = unique("INV-")
    machine = _line(client, admin, inventurni_broj=inventory)
    # An inventory number names one asset: not reusable by another machine or any component.
    assert client.post("/api/oprema", json={"naziv": unique("L "), "inventurni_broj": inventory.lower()}, headers=admin).status_code == 400
    assert client.post(f"/api/oprema/{machine['id']}/komponente", json={"naziv": "X", "inventurni_broj": inventory}, headers=admin).status_code == 400
    # A component's serial number is unique per manufacturer only.
    _component(client, admin, machine["id"], naziv="A", serijski_broj="777", proizvodac="PCMC")
    assert client.post(f"/api/oprema/{machine['id']}/komponente", json={"naziv": "B", "serijski_broj": "777", "proizvodac": "pcmc"}, headers=admin).status_code == 400
    assert client.post(f"/api/oprema/{machine['id']}/komponente", json={"naziv": "C", "serijski_broj": "777", "proizvodac": "STAX"}, headers=admin).status_code == 201
    # Same name twice on one machine is refused (ignoring case and spacing); fine on another machine.
    assert client.post(f"/api/oprema/{machine['id']}/komponente", json={"naziv": " a "}, headers=admin).status_code == 400
    assert client.post(f"/api/oprema/{_line(client, admin)['id']}/komponente", json={"naziv": "A"}, headers=admin).status_code == 201
    # Malformed values are refused, not stored.
    assert client.put(f"/api/oprema/{machine['id']}", json={"inventurni_broj": "<script>"}, headers=admin).status_code == 422


def test_component_with_faults_can_be_deactivated_not_deleted(client, admin, make_user):
    _, worker = make_user()
    machine = _line(client, admin)
    used = _component(client, admin, machine["id"])
    unused = _component(client, admin, machine["id"])
    report(client, worker, {"id": machine["id"]}, komponenta_id=used["id"])

    assert client.delete(f"/api/oprema/komponente/{used['id']}", headers=admin).status_code == 409
    assert client.put(f"/api/oprema/komponente/{used['id']}", json={"aktivna": False}, headers=admin).json()["aktivna"] is False
    assert client.delete(f"/api/oprema/komponente/{unused['id']}", headers=admin).status_code == 204
    # History stays readable on the deactivated component.
    faults = client.get(f"/api/kvarovi?komponenta={used['id']}", headers=admin).json()
    assert len(faults) == 1 and faults[0]["komponenta"]["model"] == "BS 2601"


def test_workers_read_machines_and_components_but_cannot_change_them(client, admin, make_user):
    _, worker = make_user()
    machine = _line(client, admin)
    component = _component(client, admin, machine["id"])
    assert client.get(f"/api/oprema/{machine['id']}/komponente", headers=worker).status_code == 200
    assert client.post(f"/api/oprema/{machine['id']}/komponente", json={"naziv": "Y"}, headers=worker).status_code == 403
    assert client.put(f"/api/oprema/komponente/{component['id']}", json={"naziv": "Y"}, headers=worker).status_code == 403
    assert client.get(f"/api/oprema/komponente/{component['id']}", headers=worker).status_code == 403
    assert client.get(f"/api/oprema/{machine['id']}/qr-oznake.pdf", headers=worker).status_code == 403
    assert client.get("/api/oprema/export.xlsx", headers=worker).status_code == 403
    assert client.get(f"/api/oprema/{machine['id']}", headers=worker).json()["ukupni_troskovi"] is None


# --- Faults against a machine or a component -----------------------------------------------------

def test_faults_target_the_whole_machine_or_one_of_its_components(client, admin, make_user):
    _, worker = make_user()
    machine = _line(client, admin)
    other = _line(client, admin)
    component = _component(client, admin, machine["id"])
    foreign = _component(client, admin, other["id"])

    whole = report(client, worker, {"id": machine["id"]})
    part = report(client, worker, {"id": machine["id"]}, komponenta_id=component["id"])
    assert whole["komponenta_id"] is None
    assert part["komponenta"] == {"id": component["id"], "naziv": component["naziv"], "model": "BS 2601", "proizvodac": "PCMC"}
    # A component of another machine is refused, on report and on later edits.
    bad = client.post("/api/kvarovi", json={"oprema_id": machine["id"], "komponenta_id": foreign["id"], "naslov": "x"}, headers=worker)
    assert bad.status_code == 404
    assert client.put(f"/api/kvarovi/{whole['id']}", json={"komponenta_id": foreign["id"]}, headers=admin).status_code == 404

    # Filters: the machine's list has both; "masina" only the machine-wide one; a component only its own.
    ids = lambda query: {k["id"] for k in client.get(f"/api/kvarovi?{query}", headers=admin).json()}  # noqa: E731
    assert ids(f"oprema_id={machine['id']}") == {whole["id"], part["id"]}
    assert ids(f"oprema_id={machine['id']}&komponenta=masina") == {whole["id"]}
    assert ids(f"komponenta={component['id']}") == {part["id"]}
    assert client.get("/api/kvarovi?komponenta=abc", headers=admin).status_code == 422

    # Counts on the machine include the component fault once.
    listed = next(m for m in client.get("/api/oprema", headers=admin).json() if m["id"] == machine["id"])
    assert listed["aktivni_kvarovi"] == 2 and listed["hitni_kvarovi"] == 2
    assert client.get(f"/api/oprema/komponente/{component['id']}", headers=admin).json()["aktivni_kvarovi"] == 1

    # Analytics: one machine row with both faults (one machine-wide), one component row.
    data = client.get("/api/dashboard/analitika?period=30d", headers=admin).json()
    row = next((m for m in data["masine"] if m["oprema_id"] == machine["id"]), None)
    if row:  # top 10 in a shared test database
        assert row["kvarovi"] == 2 and row["na_cijeloj_masini"] == 1

    # Export names the component (empty = whole machine).
    sheet = load_workbook(BytesIO(client.get(f"/api/kvarovi/export.xlsx?oprema_id={machine['id']}", headers=admin).content)).active
    header = [c.value for c in sheet[1]]
    rows = {r[0]: r for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert rows[part["id"]][header.index("Komponenta")] == component["naziv"]
    assert rows[whole["id"]][header.index("Komponenta")] in ("", None)


def test_inactive_machines_and_components_take_no_new_reports(client, admin, make_user):
    _, worker = make_user()
    machine = _line(client, admin)
    component = _component(client, admin, machine["id"])
    client.put(f"/api/oprema/komponente/{component['id']}", json={"aktivna": False}, headers=admin)
    refused = client.post("/api/kvarovi", json={"oprema_id": machine["id"], "komponenta_id": component["id"], "naslov": "x"}, headers=worker)
    assert refused.status_code == 400
    client.put(f"/api/oprema/{machine['id']}", json={"status": "Raspisano"}, headers=admin)
    assert client.post("/api/kvarovi", json={"oprema_id": machine["id"], "naslov": "x"}, headers=worker).status_code == 400


# --- QR -------------------------------------------------------------------------------------------

def test_qr_identifiers_resolve_to_their_machine_or_component(client, admin, make_user):
    _, worker = make_user()
    machine = _line(client, admin, proizvodac="PCMC")
    component = _component(client, admin, machine["id"])

    by_machine = client.get(f"/api/qr/{machine['qr_oznaka']}", headers=worker).json()
    assert by_machine["vrsta"] == "masina" and by_machine["oprema"]["id"] == machine["id"] and by_machine["komponenta"] is None
    by_component = client.get(f"/api/qr/{component['qr_oznaka']}", headers=worker).json()
    assert by_component["vrsta"] == "komponenta" and by_component["oprema"]["id"] == machine["id"]
    assert by_component["komponenta"]["id"] == component["id"] and by_component["komponenta"]["model"] == "BS 2601"
    assert "ukupni_troskovi" not in by_machine["oprema"] and "inventurni_broj" not in by_component["komponenta"]

    assert client.get("/api/qr/nepostojeca-oznaka", headers=worker).status_code == 404
    assert client.get("/api/qr/..%2F..%2Fetc", headers=worker).status_code == 404
    assert client.get(f"/api/qr/{machine['qr_oznaka']}").status_code == 401

    client.put(f"/api/oprema/komponente/{component['id']}", json={"aktivna": False}, headers=admin)
    assert client.get(f"/api/qr/{component['qr_oznaka']}", headers=worker).json()["komponenta"]["aktivna"] is False


def test_regenerating_a_qr_identifier_is_explicit_and_retires_the_old_one(client, admin):
    machine = _line(client, admin)
    component = _component(client, admin, machine["id"])
    renewed = client.post(f"/api/oprema/komponente/{component['id']}/nova-qr-oznaka", headers=admin).json()
    assert renewed["qr_oznaka"] != component["qr_oznaka"]
    assert client.get(f"/api/qr/{component['qr_oznaka']}", headers=admin).status_code == 404
    assert client.get(f"/api/qr/{renewed['qr_oznaka']}", headers=admin).json()["komponenta"]["id"] == component["id"]
    # Printing labels again does not change it.
    client.get(f"/api/oprema/komponente/{component['id']}/qr-oznaka.pdf", headers=admin)
    assert client.get(f"/api/oprema/komponente/{component['id']}", headers=admin).json()["qr_oznaka"] == renewed["qr_oznaka"]


def test_label_sheet_holds_the_machine_and_each_component_and_decodes_to_their_links(client, admin):
    machine = _line(client, admin)
    components = [_component(client, admin, machine["id"], naziv=f"ODMOTAČ {n}") for n in (1, 2, 3)]
    retired = _component(client, admin, machine["id"], naziv="STARI DIO")
    client.put(f"/api/oprema/komponente/{retired['id']}", json={"aktivna": False}, headers=admin)

    response = client.get(f"/api/oprema/{machine['id']}/qr-oznake.pdf", headers=admin)
    assert response.status_code == 200 and response.headers["content-type"] == "application/pdf"
    decoded = _decode_pdf(response.content)
    expected = {f"http://testserver/q/{machine['qr_oznaka']}"} | {f"http://testserver/q/{c['qr_oznaka']}" for c in components}
    assert set(decoded) == expected  # the deactivated component is left out
    # The names are printed in the text layer.
    text = "".join(page.get_textpage().get_text_range() for page in pypdfium2.PdfDocument(response.content))
    assert all(f"ODMOTAČ {n}" in text for n in (1, 2, 3)) and machine["naziv"] in text and "STARI DIO" not in text

    single = client.get(f"/api/oprema/komponente/{components[1]['id']}/qr-oznaka.pdf", headers=admin)
    assert _decode_pdf(single.content) == [f"http://testserver/q/{components[1]['qr_oznaka']}"]
    svg = client.get(f"/api/qr/{components[1]['qr_oznaka']}/slika.svg", headers=admin)
    assert svg.status_code == 200 and svg.text.lstrip().startswith("<?xml")


def test_whole_plant_machine_labels_use_the_opaque_link(client, admin):
    machine = _line(client, admin)
    decoded = _decode_pdf(client.get(f"/api/oprema/qr-oznake.pdf?ids={machine['id']}", headers=admin).content)
    assert decoded == [f"http://testserver/q/{machine['qr_oznaka']}"]


# --- Seeding ----------------------------------------------------------------------------------------

def test_seed_creates_amica_matrix_once_and_never_overwrites_entered_data(client):
    db = SessionLocal()
    try:
        dry = seed_amica_matrix(db, apply=False)
        assert not dry.applied and dry.machine in ("created", "existing")
        first = seed_amica_matrix(db, apply=True)
        assert first.applied and first.ok
        machine = db.query(Oprema).filter(Oprema.naziv == MACHINE).one()
        components = db.query(KomponentaOpreme).filter(KomponentaOpreme.oprema_id == machine.id).all()
        assert len(components) == 13
        assert {(c.naziv, c.model, c.proizvodac) for c in components} == set(COMPONENTS)
        assert all(c.inventurni_broj is None and c.serijski_broj is None and c.qr_oznaka for c in components)
        assert machine.inventurni_broj is None and machine.serijski_broj is None

        # Entered by an admin later: kept. A different model: reported, not overwritten.
        odmotac = next(c for c in components if c.naziv == "ODMOTAČ 2")
        odmotac.inventurni_broj = "INV-REAL-2"
        odmotac.model = "BS 2602-X"
        token = odmotac.qr_oznaka
        db.commit()
        again = seed_amica_matrix(db, apply=True)
        assert again.machine == "existing" and not again.created
        assert any("ODMOTAČ 2" in line for line in again.conflicts)
        db.refresh(odmotac)
        assert (odmotac.inventurni_broj, odmotac.model, odmotac.qr_oznaka) == ("INV-REAL-2", "BS 2602-X", token)
        assert db.query(KomponentaOpreme).filter(KomponentaOpreme.oprema_id == machine.id).count() == 13
        assert db.query(Oprema).filter(Oprema.naziv == MACHINE).count() == 1
    finally:
        db.close()


def test_seed_stops_when_the_machine_is_ambiguous(client):
    db = SessionLocal()
    try:
        # Two look-alikes, whether or not the real one exists already.
        db.add_all([Oprema(naziv="amica  matrix"), Oprema(naziv="Amica Matrix ")])
        db.commit()
        before = db.query(KomponentaOpreme).count()
        result = seed_amica_matrix(db, apply=True)
        assert result.machine == "ambiguous" and not result.applied
        assert db.query(KomponentaOpreme).count() == before
    finally:
        db.query(Oprema).filter(Oprema.naziv.in_(["amica  matrix", "Amica Matrix "])).delete()
        db.commit()
        db.close()


# --- Guarded test-data reset --------------------------------------------------------------------------

@pytest.fixture()
def disposable(tmp_path):
    """A separate SQLite database with users, a machine, a component, a part and two faults."""
    engine = create_engine(f"sqlite:///{tmp_path / 'test-cycle.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        user = Korisnik(username="w", ime="W", prezime="W", hashed_password="x")
        machine = Oprema(naziv="AMICA MATRIX")
        db.add_all([user, machine])
        db.flush()
        component = KomponentaOpreme(oprema_id=machine.id, naziv="ODMOTAČ 1")
        part = RezervniDio(naziv="Ležaj")
        db.add_all([component, part])
        db.flush()
        for _ in range(2):
            fault = Kvar(oprema_id=machine.id, komponenta_id=component.id, prijavio_id=user.id, naslov="test")
            db.add(fault)
            db.flush()
            db.add_all([
                KvarDogadjaj(kvar_id=fault.id, tip="created", naslov="Kvar prijavljen"),
                KvarPrilog(kvar_id=fault.id, naziv="a.jpg", mime="image/jpeg", velicina=1, sadrzaj=b"x"),
                KvarRezervniDio(kvar_id=fault.id, rezervni_dio_id=part.id, kolicina=1),
            ])
        db.commit()
    return engine


SAFE_ENV = {}


def test_reset_needs_the_marker_the_right_name_and_the_flag(disposable):
    with pytest.raises(reset_test_data.Refused):
        reset_test_data.reset_faults(disposable, "test-cycle.db", True, environ=SAFE_ENV)  # not marked
    with pytest.raises(reset_test_data.Refused):
        reset_test_data.mark_disposable(disposable, "other.db", environ=SAFE_ENV)  # wrong name
    reset_test_data.mark_disposable(disposable, "test-cycle.db", environ=SAFE_ENV)
    with pytest.raises(reset_test_data.Refused):
        reset_test_data.reset_faults(disposable, "test-cycle.db", False, environ=SAFE_ENV)  # no --yes flag

    deleted = reset_test_data.reset_faults(disposable, "test-cycle.db", True, environ=SAFE_ENV)
    assert deleted == {"kvarovi": 2, "kvar_dogadjaji": 2, "kvar_prilozi": 2, "kvar_rezervni_dijelovi": 2}
    with Session(disposable) as db:
        assert reset_test_data.counts(db) == {"kvarovi": 0, "kvar_dogadjaji": 0, "kvar_prilozi": 0, "kvar_rezervni_dijelovi": 0}
        # Users, machines, components and the catalogue stay.
        assert db.query(Korisnik).count() == 1 and db.query(Oprema).count() == 1
        assert db.query(KomponentaOpreme).count() == 1 and db.query(RezervniDio).count() == 1
    # Idempotent.
    assert reset_test_data.reset_faults(disposable, "test-cycle.db", True, environ=SAFE_ENV)["kvarovi"] == 0


@pytest.mark.parametrize("environ", [
    {"RAILWAY_ENVIRONMENT_NAME": "production"},
    {"RAILWAY_ENVIRONMENT": "Production"},
    {"APP_ENV": "production"},
])
def test_reset_is_refused_in_production_even_when_everything_else_is_set(disposable, environ):
    reset_test_data.mark_disposable(disposable, "test-cycle.db", environ=SAFE_ENV)
    with pytest.raises(reset_test_data.Refused, match="production"):
        reset_test_data.reset_faults(disposable, "test-cycle.db", True, environ=environ)
    with pytest.raises(reset_test_data.Refused, match="production"):
        reset_test_data.mark_disposable(disposable, "test-cycle.db", environ=environ)
    with Session(disposable) as db:
        assert reset_test_data.counts(db)["kvarovi"] == 2


def test_reset_refuses_a_remote_database_unless_that_host_is_named(monkeypatch, disposable):
    remote = reset_test_data.Target("postgresql", "railway", "shinkansen.proxy.rlwy.net")
    monkeypatch.setattr(reset_test_data, "describe", lambda engine: remote)
    with pytest.raises(reset_test_data.Refused, match="remote host"):
        reset_test_data.check_target(disposable, "railway", None, environ=SAFE_ENV)
    with pytest.raises(reset_test_data.Refused, match="remote host"):
        reset_test_data.check_target(disposable, "railway", "other.host", environ=SAFE_ENV)
    assert reset_test_data.check_target(disposable, "railway", "shinkansen.proxy.rlwy.net", environ=SAFE_ENV) == remote
